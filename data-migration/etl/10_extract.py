"""Stage 1 EXTRACT - copy every spreadsheet cell, untouched, into staging/raw/.

One JSONL file per workbook; one line per non-empty row:
  {"sheet": name, "row": excel_row_number, "cells": {"A": value, ...}, "f": {"C": "=SUM(...)"}}
Dates are written as "@date:<iso>" so they stay distinguishable from text.
The source archive is opened read-only. Re-running overwrites staging/raw/ only.

Usage: python data-migration/etl/10_extract.py [--dry-run]
"""
import argparse
import csv
import json
import os
import sys
from datetime import date, datetime, time

sys.path.insert(0, os.path.dirname(__file__))
from lib import DATE_PREFIX, RAW_DIR, REPO, REPORTS, col_letter, utf8_stdout  # noqa: E402

SOURCE_DIR = REPO / "05. Trust - various xl sheets"


def enc(v):
    if isinstance(v, datetime):
        return DATE_PREFIX + v.isoformat()
    if isinstance(v, date):
        return DATE_PREFIX + datetime(v.year, v.month, v.day).isoformat()
    if isinstance(v, time):
        return v.isoformat()
    if isinstance(v, float) and v != v:
        return None
    if isinstance(v, (int, float, bool, str)):
        return v
    return str(v)


def openxml(path):
    import openpyxl
    wb_v = openpyxl.load_workbook(path, data_only=True)
    wb_f = openpyxl.load_workbook(path, data_only=False)
    for ws in wb_v.worksheets:
        wf = wb_f[ws.title]
        rows = {}
        for (r, c), cell in ws._cells.items():
            if cell.value is not None and cell.value != "":
                rows.setdefault(r, {})[col_letter(c)] = enc(cell.value)
        formulas = {}
        for (r, c), cell in wf._cells.items():
            if cell.data_type == "f":
                text = getattr(cell.value, "text", cell.value)
                formulas.setdefault(r, {})[col_letter(c)] = str(text)
                rows.setdefault(r, {})      # a formula with no cached value still marks the row
        for r in sorted(rows):
            if rows[r] or formulas.get(r):
                yield ws.title, r, rows[r], formulas.get(r)


def xls(path):
    import xlrd
    wb = xlrd.open_workbook(path)
    for sh in wb.sheets():
        for r in range(sh.nrows):
            cells = {}
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                v = cell.value
                if cell.ctype in (0, 6) or v == "":
                    continue
                if cell.ctype == 3:
                    try:
                        v = xlrd.xldate_as_datetime(v, wb.datemode)
                    except Exception:
                        v = f"#BAD_DATE_SERIAL:{v}"
                elif cell.ctype == 5:
                    v = "#ERROR"
                elif cell.ctype == 4:
                    v = bool(v)
                cells[col_letter(c + 1)] = enc(v)
            if cells:
                yield sh.name, r + 1, cells, None


def csv_file(path):
    raw = open(path, "rb").read()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1252", errors="replace")
    for r, row in enumerate(csv.reader(text.splitlines()), 1):
        cells = {col_letter(c): v for c, v in enumerate(row, 1) if v.strip()}
        if cells:
            yield "(csv)", r, cells, None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    utf8_stdout()
    inv = list(csv.DictReader(open(REPORTS / "00_inventory.csv", encoding="utf-8-sig")))
    if not args.dry_run:
        RAW_DIR.mkdir(parents=True, exist_ok=True)
    files = rows = cells_n = 0
    for r in inv:
        if r["ext"] not in (".xlsx", ".xlsm", ".xls", ".csv") or r["name"].startswith("~$"):
            continue
        path = SOURCE_DIR / r["path"]
        reader = csv_file if r["ext"] == ".csv" else openxml if r["magic"] == "zip" else xls
        lines = []
        for sheet, rownum, cells, formulas in reader(path):
            rec = {"sheet": sheet, "row": rownum, "cells": cells}
            if formulas:
                rec["f"] = formulas
            lines.append(json.dumps(rec, ensure_ascii=False))
            cells_n += len(cells)
        files, rows = files + 1, rows + len(lines)
        if not args.dry_run:
            (RAW_DIR / f"{r['file_id']}.jsonl").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"workbooks extracted: {files}  rows: {rows}  cells: {cells_n}")
    print("dry-run: nothing written" if args.dry_run else f"wrote {RAW_DIR}")


if __name__ == "__main__":
    main()
