"""Phase 0.3 - profile every sheet of every workbook, and every document (read-only).

Reads reports/00_inventory.csv. Writes:
  reports/profiles/<file_id>.json   full per-sheet profile (contains sample rows = personal data)
  reports/01_sheets.csv             one line per sheet
  reports/01_documents.csv          one line per pdf/docx/other file
  reports/doc_text/<file_id>.txt    extracted text of pdf/docx files
  reports/00_duplicates_near.csv    sheet pairs sharing most of their rows
  reports/01_total_checks.csv       printed totals vs. sum of the rows above them

Usage: python data-migration/scripts/01_profile.py [--dry-run] [--only F012,F013]
"""
import argparse
import csv
import hashlib
import json
import re
import traceback
from collections import Counter
from datetime import date, datetime

from common import (REPORTS, SOURCE_DIR, cell_kind, clean_text, jsonable, script_of, utf8_stdout)

TOTAL_RE = re.compile(r"\b(grand\s*total|sub\s*-?total|total|balance|c/f|b/f|carried|brought)\b|एकूण|बेरीज|शिल्लक", re.I)
AMOUNT_RE = re.compile(r"amount|amt|rs\.?|₹|रक्कम|रुपये|donation|fees?|debit|credit|deposit|withdraw|principal|balance|देणगी", re.I)
NAME_RE = re.compile(r"\bname\b|नाव|नांव|donor|member|payee|particular", re.I)

DOMAIN_KEYWORDS = {
    "membership_register": r"member|सभासद|membership|silver|gold|platinum|renew|due date|expiry",
    "donation_register": r"donat|देणगी|donor|80g",
    "receipt_book": r"receipt|पावती|rcpt",
    "expense_ledger": r"expense|voucher|खर्च|payment|paid to|bill|invoice",
    "bank_statement": r"passbook|statement|narration|chq|cheque|withdrawal|deposit|utr|neft|imps|upi",
    "fd_register": r"\bfd\b|\bfds\b|fixed deposit|term deposit|\btd\b|\btds\b|maturity|मुदत ठेव|rate of interest|roi",
    "scheme_beneficiary": r"scholarship|dnyan|kiran|diya|niramay|udaan|udan|medical aid|student|beneficiar|applicat|विद्यार्थी",
    "event_accounts": r"gaurav|sohala|sohola|sammelan|ek divas|kayasth|event|15.?aug|svds",
    "budget": r"budget|estimate|अंदाजपत्रक",
    "audit_balance_sheet": r"balance sheet|income (and|&) expenditure|receipts (and|&) payments|audit|trial balance|utili[sz]ation|schedule",
}


# ---------- readers: each returns a list of sheet dicts with a dense `grid` ----------

def read_openxml(path):
    import openpyxl
    wb_v = openpyxl.load_workbook(path, data_only=True)
    wb_f = openpyxl.load_workbook(path, data_only=False)
    ext_links = []
    for l in getattr(wb_f, "_external_links", []):
        try:
            ext_links.append(l.file_link.Target)
        except Exception:
            ext_links.append("?")
    sheets = []
    for ws_v in wb_v.worksheets:
        ws_f = wb_f[ws_v.title]
        cells = {k: c.value for k, c in ws_v._cells.items() if c.value is not None}
        max_r = max((r for r, _ in cells), default=0)
        max_c = max((c for _, c in cells), default=0)
        grid = [[cells.get((r, c)) for c in range(1, max_c + 1)] for r in range(1, max_r + 1)]
        formulas, stale, ext_refs, samples, comments = 0, 0, 0, [], 0
        for (r, c), cell in ws_f._cells.items():
            if cell.comment is not None:
                comments += 1
            if cell.data_type == "f":
                formulas += 1
                text = getattr(cell.value, "text", cell.value)
                text = str(text)
                if "[" in text:
                    ext_refs += 1
                if ws_v._cells.get((r, c)) is None or ws_v._cells[(r, c)].value is None:
                    stale += 1
                if len(samples) < 6:
                    samples.append(f"{cell.coordinate}: {text[:120]}")
        sheets.append({
            "name": ws_v.title, "state": ws_v.sheet_state, "grid": grid,
            "merged": [str(m) for m in ws_f.merged_cells.ranges],
            "formula_count": formulas, "formulas_without_cached_value": stale,
            "formulas_referencing_other_workbooks": ext_refs, "formula_samples": samples,
            "comment_count": comments,
        })
    return sheets, {"external_links": ext_links,
                    "chartsheets": [c.title for c in wb_v.chartsheets],
                    "sheet_order": wb_v.sheetnames}


def read_xls(path):
    import xlrd
    fmt = True
    try:
        wb = xlrd.open_workbook(path, formatting_info=True)
    except Exception:
        fmt = False
        wb = xlrd.open_workbook(path)
    sheets = []
    for sh in wb.sheets():
        grid = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                v = cell.value
                if cell.ctype in (0, 6) or v == "":
                    v = None
                elif cell.ctype == 3:
                    try:
                        v = xlrd.xldate_as_datetime(v, wb.datemode)
                    except Exception:
                        v = f"#BAD_DATE_SERIAL:{v}"
                elif cell.ctype == 5:
                    v = "#ERROR"
                elif cell.ctype == 4:
                    v = bool(v)
                row.append(v)
            grid.append(row)
        sheets.append({
            "name": sh.name, "state": {0: "visible", 1: "hidden", 2: "veryHidden"}.get(sh.visibility, "?"),
            "grid": grid,
            "merged": [f"r{a + 1}-{b}:c{c + 1}-{d}" for a, b, c, d in sh.merged_cells] if fmt else None,
            "formula_count": None,  # xlrd exposes cached values only
            "formulas_without_cached_value": None, "formulas_referencing_other_workbooks": None,
            "formula_samples": [], "comment_count": len(getattr(sh, "cell_note_map", {}) or {}),
        })
    return sheets, {"external_links": None, "chartsheets": [], "sheet_order": wb.sheet_names(),
                    "note": "legacy .xls: formulas not inspectable, cached values only"}


def read_csv(path):
    raw = open(path, "rb").read()
    enc = "utf-8-sig"
    try:
        text = raw.decode(enc)
    except UnicodeDecodeError:
        enc = "cp1252"
        text = raw.decode(enc, errors="replace")
    grid = [[(v if v.strip() else None) for v in row] for row in csv.reader(text.splitlines())]
    width = max((len(r) for r in grid), default=0)
    grid = [r + [None] * (width - len(r)) for r in grid]
    return [{"name": "(csv)", "state": "visible", "grid": grid, "merged": [], "formula_count": 0,
             "formulas_without_cached_value": 0, "formulas_referencing_other_workbooks": 0,
             "formula_samples": [], "comment_count": 0}], {"encoding": enc, "external_links": [], "chartsheets": []}


# ---------- analysis ----------

def num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def norm_for_hash(v):
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return clean_text(v).lower()


def find_header(grid):
    """First row (within the top 40) that is mostly text, has >=3 cells and is followed by data."""
    best = None
    for i, row in enumerate(grid[:40]):
        vals = [v for v in row if cell_kind(v) != "empty"]
        if len(vals) < 3:
            continue
        texts = sum(1 for v in vals if cell_kind(v) == "text")
        if texts / len(vals) < 0.7:
            continue
        following = [r for r in grid[i + 1:i + 4] if sum(cell_kind(v) != "empty" for v in r) >= 2]
        if not following:
            continue
        score = len(vals)
        if best is None or score > best[1] * 1.5:  # prefer the earliest unless a later row is much wider
            best = (i, score)
    return best[0] if best else None


def analyse(sheet, file_name):
    grid = sheet["grid"]
    n_rows = len(grid)
    n_cols = max((len(r) for r in grid), default=0)
    nonempty_idx = [i for i, r in enumerate(grid) if any(cell_kind(v) != "empty" for v in r)]
    out = {k: v for k, v in sheet.items() if k != "grid"}
    out.update({"rows_used": n_rows, "cols_used": n_cols, "nonempty_rows": len(nonempty_idx),
                "empty_row_ratio": round(1 - len(nonempty_idx) / n_rows, 3) if n_rows else None,
                "merged_count": len(sheet["merged"]) if sheet["merged"] is not None else None})
    out["merged"] = (sheet["merged"] or [])[:40]
    if not nonempty_idx:
        out["empty_sheet"] = True
        return out, set()

    h = find_header(grid)
    out["header_row"] = h + 1 if h is not None else None
    out["title_rows"] = [[jsonable(v) for v in grid[i] if v is not None][:6] for i in nonempty_idx if h is not None and i < h][:8]
    data_start = (h + 1) if h is not None else 0
    header = grid[h] if h is not None else [None] * n_cols

    total_rows = []
    for i in nonempty_idx:
        if i < data_start:
            continue
        if any(isinstance(v, str) and TOTAL_RE.search(v) for v in grid[i]):
            total_rows.append(i)
    out["total_rows"] = [i + 1 for i in total_rows][:60]
    total_set = set(total_rows)
    data_idx = [i for i in nonempty_idx if i >= data_start and i not in total_set]
    out["data_rows"] = len(data_idx)

    columns, flags = [], Counter()
    for c in range(n_cols):
        kinds = Counter(cell_kind(grid[i][c]) for i in data_idx if c < len(grid[i]))
        filled = sum(v for k, v in kinds.items() if k != "empty")
        hname = clean_text(header[c]) if c < len(header) and header[c] is not None else ""
        if not filled and not hname:
            continue
        col = {"col": c + 1, "header": hname, "header_script": script_of(hname) if hname else None,
               "filled": filled, "kinds": dict(kinds)}
        numeric = [num(grid[i][c]) for i in data_idx if c < len(grid[i]) and num(grid[i][c]) is not None]
        if numeric:
            col.update({"sum": round(sum(numeric), 2), "min": min(numeric), "max": max(numeric),
                        "non_integer": sum(1 for x in numeric if not float(x).is_integer())})
        dates = [grid[i][c] for i in data_idx if c < len(grid[i]) and isinstance(grid[i][c], (datetime, date))]
        if dates:
            ds = [d if isinstance(d, datetime) else datetime(d.year, d.month, d.day) for d in dates]
            col.update({"date_min": min(ds).date().isoformat(), "date_max": max(ds).date().isoformat()})
            bad = sum(1 for d in ds if d.year < 2015 or d > datetime(2027, 3, 31))
            if bad:
                col["dates_outside_2015_to_FY2026_27"] = bad
                flags["implausible_dates"] += bad
        texts = [clean_text(grid[i][c]) for i in data_idx if c < len(grid[i]) and cell_kind(grid[i][c]) == "text"]
        if texts:
            col["text_scripts"] = dict(Counter(script_of(t) for t in texts))
            col["distinct_text"] = len(set(texts))
        if filled and kinds.get("number", 0) / filled >= 0.8 and (kinds.get("text", 0) or kinds.get("number_as_text", 0)):
            col["text_in_numeric_column"] = kinds.get("text", 0) + kinds.get("number_as_text", 0)
            flags["text_in_numeric_column"] += col["text_in_numeric_column"]
        if kinds.get("number_as_text"):
            flags["numbers_stored_as_text"] += kinds["number_as_text"]
        if kinds.get("date_as_text"):
            flags["dates_stored_as_text"] += kinds["date_as_text"]
        if AMOUNT_RE.search(hname) and numeric:
            neg, zero = sum(1 for x in numeric if x < 0), sum(1 for x in numeric if x == 0)
            if neg:
                col["negative_amounts"] = neg
                flags["negative_amounts"] += neg
            if zero:
                col["zero_amounts"] = zero
                flags["zero_amounts"] += zero
        if NAME_RE.search(hname) and len(data_idx) > 3:
            blanks = kinds.get("empty", 0)
            if blanks:
                col["blank_in_name_column"] = blanks
                flags["blank_names"] += blanks
        columns.append(col)
    out["columns"] = columns
    out["quality_flags"] = dict(flags)

    # printed totals vs. the rows they claim to add up
    checks, prev_boundary = [], data_start - 1
    for t in total_rows:
        for c in range(n_cols):
            printed = num(grid[t][c]) if c < len(grid[t]) else None
            if printed is None or printed == 0:
                continue
            seg = [num(grid[i][c]) for i in range(prev_boundary + 1, t) if c < len(grid[i]) and i not in total_set]
            seg_sum = sum(x for x in seg if x is not None)
            all_sum = sum(x for i in data_idx if i < t and c < len(grid[i]) for x in [num(grid[i][c])] if x is not None)
            sub_sum = sum(x for i in total_rows if i < t and c < len(grid[i]) for x in [num(grid[i][c])] if x is not None)
            if not any(x is not None for x in seg) and not all_sum:
                continue
            verdict = ("match_segment" if abs(seg_sum - printed) < 0.005 else
                       "match_all_rows_above" if abs(all_sum - printed) < 0.005 else
                       "match_sum_of_subtotals" if sub_sum and abs(sub_sum - printed) < 0.005 else "UNEXPLAINED")
            checks.append({"total_row": t + 1, "col": c + 1,
                           "header": clean_text(header[c]) if c < len(header) and header[c] is not None else "",
                           "label": next((clean_text(v)[:40] for v in grid[t] if isinstance(v, str) and TOTAL_RE.search(v)), ""),
                           "printed": printed, "segment_sum": round(seg_sum, 2), "all_rows_sum": round(all_sum, 2),
                           "verdict": verdict, "diff_vs_segment": round(printed - seg_sum, 2)})
        prev_boundary = t
    out["total_checks"] = checks

    out["sample_head"] = [[jsonable(v) for v in grid[i]] for i in data_idx[:6]]
    out["sample_tail"] = [[jsonable(v) for v in grid[i]] for i in (data_idx[-3:] if len(data_idx) > 9 else [])]
    out["total_row_samples"] = [[jsonable(v) for v in grid[i]] for i in total_rows[:4]]

    hay = " ".join([file_name, sheet["name"]] + [c["header"] for c in columns] +
                   [str(x) for t in out["title_rows"] for x in t]).lower()
    hits = {d: len(re.findall(p, hay)) for d, p in DOMAIN_KEYWORDS.items()}
    out["domain_keyword_hits"] = {d: n for d, n in hits.items() if n}

    fp = set()
    for i in data_idx:
        vals = [norm_for_hash(v) for v in grid[i] if cell_kind(v) != "empty"]
        if len(vals) >= 2:
            fp.add(hashlib.sha1("|".join(vals).encode("utf-8")).hexdigest()[:16])
    out["distinct_row_fingerprints"] = len(fp)
    return out, fp


def profile_document(path, row):
    info = {"file_id": row["file_id"], "path": row["path"], "ext": row["ext"], "pages": "", "text_chars": 0,
            "tables": "", "likely_scanned": "", "error": "", "preview": ""}
    text = ""
    try:
        if row["magic"] == "pdf":
            from pypdf import PdfReader
            rd = PdfReader(str(path))
            info["pages"] = len(rd.pages)
            text = "\n\n".join(f"--- page {n} ---\n{(p.extract_text() or '')}" for n, p in enumerate(rd.pages, 1))
            body = re.sub(r"--- page \d+ ---", "", text)
            info["likely_scanned"] = len(body.strip()) < 40 * len(rd.pages)
        elif row["ext"] == ".docx":
            import docx
            d = docx.Document(str(path))
            parts = [p.text for p in d.paragraphs if p.text.strip()]
            for ti, t in enumerate(d.tables, 1):
                parts.append(f"--- table {ti} ---")
                parts += [" | ".join(c.text.strip() for c in r.cells) for r in t.rows]
            info["tables"] = len(d.tables)
            text = "\n".join(parts)
        elif row["ext"] in (".txt", ".bas"):
            text = open(path, "rb").read().decode("utf-8", errors="replace")
        else:
            info["error"] = "not text-extractable here (image / legacy .doc / pptx)"
    except Exception as e:
        info["error"] = f"{type(e).__name__}: {e}"[:200]
    text = clean_nfc_block(text)
    info["text_chars"] = len(text)
    info["preview"] = clean_text(text)[:240]
    return info, text


def clean_nfc_block(text):
    import unicodedata
    return unicodedata.normalize("NFC", text)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--only", default="")
    args = ap.parse_args()
    utf8_stdout()

    inv = list(csv.DictReader(open(REPORTS / "00_inventory.csv", encoding="utf-8-sig")))
    only = set(filter(None, args.only.split(",")))
    seen_hash, sheet_rows, doc_rows, check_rows, fps, errors = {}, [], [], [], {}, []
    if not args.dry_run:
        (REPORTS / "profiles").mkdir(parents=True, exist_ok=True)
        (REPORTS / "doc_text").mkdir(parents=True, exist_ok=True)

    for row in inv:
        if only and row["file_id"] not in only:
            continue
        path = SOURCE_DIR / row["path"]
        is_sheet = row["ext"] in (".xlsx", ".xlsm", ".xls", ".csv")
        if row["sha256"] in seen_hash:  # byte-identical copy: profile once
            if is_sheet:
                sheet_rows.append({"file_id": row["file_id"], "path": row["path"], "sheet": "*",
                                   "note": f"exact duplicate of {seen_hash[row['sha256']]}"})
            else:
                doc_rows.append({"file_id": row["file_id"], "path": row["path"], "ext": row["ext"],
                                 "error": f"exact duplicate of {seen_hash[row['sha256']]}"})
            continue
        seen_hash[row["sha256"]] = row["file_id"]

        if not is_sheet:
            info, text = profile_document(path, row)
            doc_rows.append(info)
            if text and not args.dry_run:
                (REPORTS / "doc_text" / f"{row['file_id']}.txt").write_text(text, encoding="utf-8")
            continue

        try:
            if row["ext"] == ".csv":
                sheets, wbmeta = read_csv(path)
            elif row["magic"] == "zip":
                sheets, wbmeta = read_openxml(path)
            elif row["magic"] == "ole2":
                sheets, wbmeta = read_xls(path)
            else:
                raise ValueError(f"unrecognised spreadsheet container ({row['magic']})")
        except Exception as e:
            errors.append((row["file_id"], row["path"], f"{type(e).__name__}: {e}"))
            sheet_rows.append({"file_id": row["file_id"], "path": row["path"], "sheet": "*",
                               "note": f"READ ERROR {type(e).__name__}: {e}"[:200]})
            traceback.print_exc()
            continue

        profile = {"file_id": row["file_id"], "path": row["path"], "sha256": row["sha256"],
                   "modified": row["modified"], "container": row["magic"], "workbook": wbmeta, "sheets": []}
        for n, sh in enumerate(sheets, 1):
            res, fp = analyse(sh, row["name"])
            res["sheet_id"] = f"{row['file_id']}.S{n:02d}"
            profile["sheets"].append(res)
            fps[res["sheet_id"]] = (row["path"], sh["name"], fp)
            hits = res.get("domain_keyword_hits", {})
            unexplained = sum(1 for c in res.get("total_checks", []) if c["verdict"] == "UNEXPLAINED")
            sheet_rows.append({
                "file_id": row["file_id"], "sheet_id": res["sheet_id"], "path": row["path"], "sheet": sh["name"],
                "state": res["state"], "rows_used": res["rows_used"], "cols_used": res["cols_used"],
                "data_rows": res.get("data_rows", 0), "header_row": res.get("header_row"),
                "empty_row_ratio": res["empty_row_ratio"], "merged_count": res["merged_count"],
                "formula_count": res["formula_count"], "formulas_no_cache": res["formulas_without_cached_value"],
                "ext_workbook_refs": res["formulas_referencing_other_workbooks"],
                "total_rows": len(res.get("total_rows", [])), "totals_unexplained": unexplained,
                "comments": res["comment_count"],
                "keyword_guess": max(hits, key=hits.get) if hits else "",
                "quality_flags": json.dumps(res.get("quality_flags", {}), ensure_ascii=False),
                "headers": " | ".join(c["header"] for c in res.get("columns", []) if c["header"])[:400],
                "note": "EMPTY" if res.get("empty_sheet") else "",
            })
            for c in res.get("total_checks", []):
                check_rows.append({"sheet_id": res["sheet_id"], "path": row["path"], "sheet": sh["name"], **c})
        if not args.dry_run:
            (REPORTS / "profiles" / f"{row['file_id']}.json").write_text(
                json.dumps(profile, ensure_ascii=False, indent=1, default=str), encoding="utf-8")

    # near-duplicate sheets: share most of their distinct rows
    near = []
    ids = [k for k, v in fps.items() if len(v[2]) >= 5]
    for a_i, a in enumerate(ids):
        fa = fps[a][2]
        for b in ids[a_i + 1:]:
            fb = fps[b][2]
            inter = len(fa & fb)
            if not inter:
                continue
            jac, cont = inter / len(fa | fb), inter / min(len(fa), len(fb))
            if jac >= 0.5 or cont >= 0.8:
                near.append({"sheet_a": a, "path_a": fps[a][0], "name_a": fps[a][1], "rows_a": len(fa),
                             "sheet_b": b, "path_b": fps[b][0], "name_b": fps[b][1], "rows_b": len(fb),
                             "shared_rows": inter, "jaccard": round(jac, 3), "containment": round(cont, 3),
                             "same_file": a.split(".")[0] == b.split(".")[0]})
    near.sort(key=lambda r: (-r["jaccard"], -r["shared_rows"]))

    n_sheets = sum(1 for r in sheet_rows if r.get("sheet") != "*")
    print(f"workbooks profiled: {len({r['file_id'] for r in sheet_rows if r.get('sheet') != '*'})}  sheets: {n_sheets}")
    print(f"documents profiled: {len(doc_rows)}  read errors: {len(errors)}")
    for e in errors:
        print("  ERROR", e)
    print(f"near-duplicate sheet pairs: {len(near)}")
    print("total checks:", dict(Counter(c["verdict"] for c in check_rows)))

    if args.dry_run:
        print("dry-run: nothing written")
        return

    def dump(name, rows, cols):
        with open(REPORTS / name, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
            w.writeheader()
            w.writerows(rows)

    dump("01_sheets.csv", sheet_rows, ["file_id", "sheet_id", "path", "sheet", "state", "rows_used", "cols_used",
                                       "data_rows", "header_row", "empty_row_ratio", "merged_count", "formula_count",
                                       "formulas_no_cache", "ext_workbook_refs", "total_rows", "totals_unexplained",
                                       "comments", "keyword_guess", "quality_flags", "headers", "note"])
    dump("01_documents.csv", doc_rows, ["file_id", "path", "ext", "pages", "text_chars", "tables", "likely_scanned",
                                        "error", "preview"])
    dump("00_duplicates_near.csv", near, ["sheet_a", "path_a", "name_a", "rows_a", "sheet_b", "path_b", "name_b",
                                          "rows_b", "shared_rows", "jaccard", "containment", "same_file"])
    dump("01_total_checks.csv", check_rows, ["sheet_id", "path", "sheet", "total_row", "col", "header", "label",
                                             "printed", "segment_sum", "all_rows_sum", "diff_vs_segment", "verdict"])
    print("wrote 01_sheets.csv, 01_documents.csv, 00_duplicates_near.csv, 01_total_checks.csv, profiles/, doc_text/")


if __name__ == "__main__":
    main()
